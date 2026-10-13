"""Regression tests for pasted HTML fragments in Word and PowerPoint exports."""

# Django Imports
from django.test import SimpleTestCase

# 3rd Party Libraries
import docx
import pptx

# Ghostwriter Libraries
from ghostwriter.modules.reportwriter import prepare_jinja2_env
from ghostwriter.modules.reportwriter.base.html_rich_text import rich_text_template
from ghostwriter.modules.reportwriter.richtext.docx import HtmlToDocx
from ghostwriter.modules.reportwriter.richtext.pptx import HtmlToPptx


class RichTextNormalizationTests(SimpleTestCase):
    """Exercise normalization through both document converters."""

    def convert(self, html):
        document = docx.Document()
        HtmlToDocx.run(html, document, None)
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        shape = slide.shapes.placeholders[1]
        shape.text_frame.clear()
        HtmlToPptx.run(html, slide, shape)
        HtmlToPptx.delete_extra_paragraph(shape)
        return document, slide, shape

    def assert_paragraphs(self, html, expected):
        document, slide, shape = self.convert(html)
        self.assertEqual([p.text for p in document.paragraphs], expected)
        # python-pptx represents an explicit line break as a vertical tab.
        self.assertEqual(
            [p.text.replace("\v", "\n") for p in shape.text_frame.paragraphs], expected
        )
        return document, slide, shape

    def test_text_fragments_preserve_paragraph_boundaries(self):
        cases = (
            ("Pasted text", ["Pasted text"]),
            ("<p>Before</p>Pasted text", ["Before", "Pasted text"]),
            ("Pasted text<p>After</p>", ["Pasted text", "After"]),
            (
                "<p>Before</p>Pasted text<p>After</p>",
                ["Before", "Pasted text", "After"],
            ),
            ("<div>First <b>bold</b></div><div>Second</div>", ["First bold", "Second"]),
            ("<div>Before<div>Nested</div>After</div>", ["Before", "Nested", "After"]),
            (
                "<div><p>Before</p>Pasted text<p>After</p></div>",
                ["Before", "Pasted text", "After"],
            ),
            ("<div><ul><li>First</li><li>Second</li></ul></div>", ["First", "Second"]),
        )
        for html, expected in cases:
            with self.subTest(html=html):
                self.assert_paragraphs(html, expected)

    def test_adjacent_inline_elements_keep_one_paragraph_and_formatting(self):
        document, _, shape = self.assert_paragraphs(
            "<p>Before</p>Pasted <b>bold</b> and "
            '<span class="italic underline">styled</span> text<br>next line'
            "<p>After</p>",
            ["Before", "Pasted bold and styled text\nnext line", "After"],
        )
        word_runs = {run.text: run for run in document.paragraphs[1].runs}
        slide_runs = {run.text: run for run in shape.text_frame.paragraphs[1].runs}
        self.assertTrue(word_runs["bold"].bold)
        self.assertTrue(slide_runs["bold"].font.bold)
        self.assertTrue(word_runs["styled"].italic)
        self.assertTrue(slide_runs["styled"].font.italic)
        self.assertTrue(word_runs["styled"].underline)
        self.assertTrue(slide_runs["styled"].font.underline)

    def test_inline_style_attributes_survive_wrapping(self):
        document, _, shape = self.assert_paragraphs(
            '<span class="bold" style="color: #50b071; font-family: Arial">Pasted text</span>',
            ["Pasted text"],
        )
        word_font = document.paragraphs[0].runs[0].font
        slide_font = shape.text_frame.paragraphs[0].runs[0].font
        self.assertEqual(str(word_font.color.rgb), "50B071")
        self.assertEqual(str(slide_font.color.rgb), "50B071")
        self.assertEqual(word_font.name, "Arial")
        self.assertEqual(slide_font.name, "Arial")
        self.assertTrue(word_font.bold)
        self.assertTrue(slide_font.bold)

    def test_comments_and_block_whitespace_are_not_exported(self):
        self.assert_paragraphs(
            "<!-- clipboard -->\n<p>Before</p>\n\t<!-- gap -->"
            "\nPasted <!-- inline -->text\n<p>After</p>\n<p></p>",
            ["Before", "Pasted text", "After"],
        )

    def test_div_inside_table_cell_stays_in_the_cell(self):
        document, slide, _ = self.convert(
            "<table><tr><td>Before<div>Nested <b>bold</b></div>After</td></tr></table>"
        )
        self.assertEqual(
            document.tables[0].cell(0, 0).text, "Before\nNested bold\nAfter"
        )
        self.assertEqual(document.paragraphs, [])
        table = next(shape.table for shape in slide.shapes if shape.has_table)
        self.assertEqual(table.cell(0, 0).text, "Before\nNested bold\nAfter")

    def test_mixed_table_cell_blocks_preserve_text_order(self):
        cases = (
            ("<td>Before<p>Middle</p>After</td>", "Before\nMiddle\nAfter"),
            (
                "<td><div>Before<blockquote>Quoted</blockquote>After</div></td>",
                "Before\nQuoted\nAfter",
            ),
        )
        for cell, expected in cases:
            with self.subTest(cell=cell):
                document, slide, _ = self.convert(f"<table><tr>{cell}</tr></table>")
                self.assertEqual(document.tables[0].cell(0, 0).text, expected)
                table = next(shape.table for shape in slide.shapes if shape.has_table)
                self.assertEqual(table.cell(0, 0).text.replace("\v", "\n"), expected)

    def test_nested_blocks_preserve_text_order(self):
        for html in (
            "<blockquote>Before<div>Nested</div>After</blockquote>",
            "<blockquote>Before<p>Nested</p>After</blockquote>",
            "<ul><li>Before<div>Nested</div>After</li></ul>",
            "<ul><li>Before<p>Nested</p>After</li></ul>",
        ):
            with self.subTest(html=html):
                document, _, shape = self.convert(html)
                self.assertEqual(
                    "\n".join(p.text for p in document.paragraphs),
                    "Before\nNested\nAfter",
                )
                self.assertEqual(
                    "\n".join(p.text for p in shape.text_frame.paragraphs),
                    "Before\nNested\nAfter",
                )

    def test_collaborative_table_wrapper_preserves_the_table(self):
        document, slide, _ = self.assert_paragraphs(
            'Before<div class="collab-table-wrapper">'
            '<div class="collab-table-caption" data-bookmark="table-one">'
            '<div class="collab-table-caption-content">Table caption</div></div>'
            "<table><tr><td><p>Cell</p></td></tr></table></div>After",
            ["Before", "After"],
        )
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(0, 0).text, "Cell")
        tables = [shape.table for shape in slide.shapes if shape.has_table]
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].cell(0, 0).text, "Cell")

    def test_page_break_markers_keep_their_block_boundaries(self):
        for marker in (
            '<br data-gw-pagebreak="true">',
            '<div class="page-break"></div>',
        ):
            with self.subTest(marker=marker):
                document, _, shape = self.convert(f"Before{marker}After")
                self.assertEqual(
                    [p.text for p in document.paragraphs], ["Before", "", "After"]
                )
                self.assertTrue(
                    document.paragraphs[1]._p.xpath('.//w:br[@w:type="page"]')
                )
                self.assertEqual(
                    [p.text for p in shape.text_frame.paragraphs], ["Before", "After"]
                )

    def test_unsafe_tags_and_links_remain_unsupported(self):
        document, _, shape = self.assert_paragraphs(
            "<div><script>Unwanted script text</script>"
            '<a href="javascript:alert(1)">Visible text</a></div>',
            ["Visible text"],
        )
        self.assertFalse(document.paragraphs[0]._p.xpath(".//w:hyperlink"))
        self.assertIsNone(shape.text_frame.paragraphs[0].runs[0].hyperlink.address)

    def test_text_generated_by_jinja_is_normalized_after_rendering(self):
        env = prepare_jinja2_env()
        rendered = rich_text_template(
            env, "<p>Before</p>{{ narrative }}<p>After</p>"
        ).render(narrative="Generated text")
        self.assert_paragraphs(rendered, ["Before", "Generated text", "After"])
